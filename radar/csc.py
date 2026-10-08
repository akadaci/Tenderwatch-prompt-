"""Lecture des cahiers des charges (documents publics des marchés sur e-Procurement).
Pour chaque appel d'offres de travaux encore ouvert : auteur de projet et passages sur les fixations,
RECOPIÉS MOT POUR MOT, avec leur emplacement (document, page, section de la table des matières).
Les familles Fischer proposées viennent du kit de prescription de Kad et sont à valider."""
import asyncio, io, json, re, zipfile
from datetime import date
import httpx
from radar.common import DATA

OUT = DATA / "csc.json"
MAX_TENDERS = 10
MAX_FILE = 60_000_000
MAX_TOTAL = 200_000_000
PRIORITE = {"Wallonie": 0, "Grand-Duché": 1, "Bruxelles": 2, "Flandre": 3}
WALLONIE = {"Liège", "Namur", "Hainaut", "Brabant wallon", "Luxembourg (prov.)"}

INCLUDE = re.compile(r"(csc|\bcs\b|_cs[_.]|cahier|charges|clauses|\bct\b|_ct[_.]|technique|bestek|technisch|lastenboek|"
                     r"m[ée]tr[ée]|meetstaat|stabilit|architect|gros.?oeuvre|ruwbouw|fa[cç]ade|gevel|hvac|[ée]lectri|lot)", re.I)
EXCLUDE = re.compile(r"(espd|dume|\buea\b|pss|s[ée]curit[ée]|veiligheid|(?:^|[\s/_.›-])plans?(?=[\s_.-]|$)|_pl_|\.dwg|\.dxf|\.jpe?g|\.png|photo|foto|"
                     r"formulaire|inschrijvingsformulier|attestation|modele.?d.?offre|offerteformulier)", re.I)
DOC_EXT = (".pdf", ".docx", ".xlsx", ".zip")
AUTEUR = re.compile(r"(auteur\s+d[eu]\s+projet|auteur\s+du\s+cahier|bureau\s+d['’]?\s*[ée]tudes?|ontwerper|studiebureau|"
                    r"architecte\s*(?:-|:)|ing[ée]nieur\s+en\s+stabilit[ée]|ing[ée]nieur\s+stabilit[ée])", re.I)
FIX = re.compile(r"\b(chevilles?|ancrages?|scellements?\s+chimiques?|scellement|goujons?|tiges?\s+filet[ée]es?|r[ée]sine\s+d['’]ancrage|"
                 r"fixations?\s+chimiques?|chevilles?\s+d['’]isolant|ETICS|coupe-feu|r[ée]sistant\s+au\s+feu|rails?\s+d['’]ancrage|"
                 r"pluggen|ankers?|verankering|draadstang|keilbouten|brandwerend|isolatiepluggen|"
                 r"ETA|ATE|[ÉE]valuation\s+Technique\s+Europ[ée]enne|NBN\s+EN\s+1992-4)\b", re.I)
MARQUES = re.compile(r"\b(hilti|w[üu]rth|fischer|spit|simpson|halfen|rawlplug|mungo|ejot|heco|sika|ancon|jordahl|pfeifer)\b", re.I)
HEADING = re.compile(r"^\s*((?:art(?:icle)?\.?\s*)?\d{1,3}(?:\.\d{1,3}){0,5}\.?|[A-Z]\d{1,2}(?:\.\d+)*|chapitre\s+\w+|titre\s+\w+|hoofdstuk\s+\w+)\s+\S.{2,90}$", re.I)
FAMILLES = [
    (r"chimique|r[ée]sine|scellement|chemisch|injectie", "Chimique FIS EM Plus ou FIS V Plus (avec tiges ou armatures)"),
    (r"isolant|etics|termoz|isolatie", "Chevilles d'isolant Termoz"),
    (r"coupe-feu|r[ée]sistant au feu|brandwerend|\bEI\s?\d", "Gamme coupe-feu fischer"),
    (r"goujon|expansion|m[ée]canique|keilbout|ancrage lourd|cheville acier", "Goujons FAZ II Plus, chevilles lourdes FH II"),
    (r"rail|support|console|draagsyst", "Systèmes de rails et supports fischer"),
    (r"cheville|plug|ma[cç]onnerie|brique", "Chevilles DuoPower ; FIS V Plus avec tamis en brique creuse"),
    (r"vis\b|charpente|ossature bois|schroef", "Vis Power-Fast II"),
]


def suggestion(text):
    for rx, fam in FAMILLES:
        if re.search(rx, text, re.I):
            return fam
    return ""


def pdf_pages(data):
    import pymupdf
    with pymupdf.open(stream=data, filetype="pdf") as d:
        return [d[i].get_text() for i in range(d.page_count)], d.get_toc()


def xlsx_pages(data):
    """Métré / meetstaat : une « page » par onglet, une ligne par rangée."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    pages, toc = [], []
    for i, ws in enumerate(wb.worksheets, 1):
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c not in (None, "")]
            if cells:
                rows.append(" | ".join(cells))
            if len(rows) > 20000:
                break
        pages.append("\n".join(rows)); toc.append([1, f"Onglet « {ws.title} »", i])
    return pages, toc


def docx_pages(data):
    import docx
    doc = docx.Document(io.BytesIO(data))
    toc, text = [], []
    for p in doc.paragraphs:
        if p.style is not None and p.style.name.lower().startswith(("heading", "titre", "kop")) and p.text.strip():
            toc.append([1, p.text.strip(), 1])
        text.append(p.text)
    return ["\n".join(text)], toc


def section_for(toc, page_no, lines, idx):
    """Titre numéroté le plus proche AU-DESSUS du passage, sinon dernière entrée de table des matières <= page."""
    for j in range(idx, max(-1, idx - 120), -1):
        l = " ".join(lines[j].split())
        if HEADING.match(l) and not FIX.search(l[:6]):
            return l[:110]
    cands = [t for t in toc if t[2] <= page_no]
    return cands[-1][1][:110] if cands else ""


def bloc_at(lines, i, n=3, maxlen=600):
    """La ligne trouvée + les suivantes du même paragraphe (arrêt au titre suivant)."""
    out = []
    for j in range(i, min(len(lines), i + n)):
        l = " ".join(lines[j].split())
        if not l:
            continue
        if j > i and HEADING.match(l):
            break
        out.append(l)
    return " ".join(out)[:maxlen]


def scan(name, pages, toc):
    auteur, fix = [], []
    for pi, text in enumerate(pages, 1):
        lines = text.splitlines()
        for i, raw in enumerate(lines):
            l = " ".join(raw.split())
            if len(l) < 6:
                continue
            ou = {"document": name, "page": pi if len(pages) > 1 else None, "section": ""}
            if AUTEUR.search(l) and len(auteur) < 6:
                bloc = bloc_at(lines, i, 3, 400)
                ou["section"] = section_for(toc, pi, lines, i)
                auteur.append({"texte": bloc, "ou": ou})
            if FIX.search(l):
                bloc = bloc_at(lines, i, 3, 600)
                score = 1 + 3 * bool(MARQUES.search(bloc)) + 2 * bool(re.search(r"\b(ETA|ATE|1992-4)\b", bloc))
                ou["section"] = section_for(toc, pi, lines, i)
                m = MARQUES.search(bloc)
                fix.append({"texte": bloc, "ou": ou, "score": score, "marque": m.group(0) if m else "",
                            "suggestion": suggestion(bloc)})
    return auteur, fix


def extract(name, data, depth=0):
    """[(nom, pages, toc)] ; ZIP ouvert récursivement en ne gardant que les documents utiles."""
    low = name.lower()
    try:
        if low.endswith(".zip") and depth < 3:
            out = []
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                for n in z.namelist():
                    base = n.rsplit("/", 1)[-1]
                    if n.endswith("/") or z.getinfo(n).file_size > MAX_FILE or EXCLUDE.search(base):
                        continue
                    if base.lower().endswith(DOC_EXT):
                        out += extract(f"{name} › {base}", z.read(n), depth + 1)
            return out
        if low.endswith(".pdf"):
            pages, toc = pdf_pages(data)
            return [(name, pages, toc)]
        if low.endswith(".docx"):
            pages, toc = docx_pages(data)
            return [(name, pages, toc)]
        if low.endswith(".xlsx"):
            pages, toc = xlsx_pages(data)
            return [(name, pages, toc)]
    except Exception as e:
        return [(name, None, f"{type(e).__name__}: {e}")]
    return []


def dedup(items, key="texte", n=25):
    seen, out = set(), []
    for it in sorted(items, key=lambda x: -x.get("score", 0)):
        k = re.sub(r"\W+", "", it[key].lower())[:160]
        if k in seen:
            continue
        seen.add(k)
        out.append({k2: v for k2, v in it.items() if k2 != "score"})
        if len(out) >= n:
            break
    return out


def find_url(o):
    """Le lien signé, où qu'il soit dans la réponse (texte, objet ou liste)."""
    if isinstance(o, str):
        return o if o.startswith("http") else ""
    if isinstance(o, dict):
        for v in o.values():
            u = find_url(v)
            if u:
                return u
    if isinstance(o, list):
        for v in o:
            u = find_url(v)
            if u:
                return u
    return ""


async def read_tender(s, wid):
    docs = await s.call(f"/api/dos/publication-workspaces/{wid}/documents?full=false&type=WORKSPACE&type=ESPD_REQUEST&type=SDI")
    lus, ignores, auteur, fix, total = [], [], [], [], 0
    async with httpx.AsyncClient(timeout=180, follow_redirects=True) as http:
        for d in docs or []:
            v = (d.get("versions") or [{}])[-1]
            name = ((v.get("document") or {}).get("originalFileName") or "").strip()
            if not name or EXCLUDE.search(name) or not name.lower().endswith(DOC_EXT):
                ignores.append(name); continue
            info = await s.call(f"/api/dos/publication-workspace-document-versions/{v['id']}/download-url?unpublished=false")
            url = find_url(info)
            if not url:
                ignores.append(f"{name} (lien absent : {str(info)[:120]})"); continue
            r = await http.get(url)
            if r.status_code != 200 or len(r.content) > MAX_FILE:
                ignores.append(f"{name} (HTTP {r.status_code}, {len(r.content)} o)"); continue
            total += len(r.content)
            for fname, pages, toc in extract(name, r.content):
                if pages is None:
                    ignores.append(f"{fname} ({toc})"); continue
                lus.append({"document": fname, "pages": len(pages), "table_des_matieres": bool(toc)})
                a, f = scan(fname, pages, toc if isinstance(toc, list) else [])
                auteur += a; fix += f
            if total > MAX_TOTAL:
                ignores.append("… arrêt : volume maximal atteint"); break
    return {"lu_le": date.today().isoformat(), "documents_lus": lus, "documents_ignores": ignores[:40],
            "auteur": dedup(auteur, n=6), "fixations": dedup(fix, n=25)}


def targets(limit=MAX_TENDERS):
    done = json.loads(OUT.read_text()) if OUT.exists() else {}
    bda = json.loads((DATA / "bda.json").read_text())["records"]
    det = json.loads((DATA / "bda_detail.json").read_text()) if (DATA / "bda_detail.json").exists() else {}
    today = date.today().isoformat()
    c = []
    for r in bda:
        wid = r["id"][4:]
        lim = (det.get(wid) or {}).get("date_limite") or ""
        prev = done.get(wid)
        retry = prev is not None and not prev.get("documents_lus") and prev.get("lu_le") != today
        if r["etape"] != 4 or (prev is not None and not retry) or not (r.get("cpv") or "").startswith(("45", "44")) or (lim and lim < today):
            continue
        prov = r.get("province") or ""
        reg = "Wallonie" if prov in WALLONIE else prov if prov in ("Bruxelles", "Grand-Duché") else "Flandre"
        c.append((PRIORITE.get(reg, 4), r.get("date_publication", ""), wid))
    c.sort(key=lambda x: (x[0], -int(x[1].replace("-", "") or 0)))
    return done, [w for _, _, w in c[:limit]]


def run(limit=MAX_TENDERS):
    from radar.bosa import Session
    done, todo = targets(limit)

    async def go():
        async with Session(pause_ms=800) as s:
            for wid in todo:
                try:
                    done[wid] = await read_tender(s, wid)
                except Exception as e:
                    done[wid] = {"lu_le": date.today().isoformat(), "erreur": f"{type(e).__name__}: {e}"[:300]}
                OUT.write_text(json.dumps(done, ensure_ascii=False, separators=(",", ":")))
    if todo:
        asyncio.run(go())
    print(f"[csc] {len(todo)} cahiers des charges lus ({len(done)} au total)")
    return len(todo)


if __name__ == "__main__":
    import sys
    run(int(sys.argv[1]) if len(sys.argv) > 1 else MAX_TENDERS)
